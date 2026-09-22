"""Tests de la importacion de destinatarios desde archivos."""

import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

from crearborradores.importers import (  # noqa: E402
    FileImportError,
    import_addresses,
    read_table,
)

try:
    import openpyxl
except ModuleNotFoundError:  # pragma: no cover
    openpyxl = None


class TestImportadores(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _write(self, name, content, encoding="utf-8"):
        path = self.dir / name
        path.write_text(content, encoding=encoding)
        return path

    def test_txt(self):
        path = self._write("lista.txt", "ana@x.com\nluis@y.com\n")
        self.assertEqual(import_addresses(path), ["ana@x.com", "luis@y.com"])

    def test_csv_con_encabezado_y_columnas(self):
        path = self._write("lista.csv", "Nombre;Correo;Area\nAna;ana@x.com;Ventas\nLuis;luis@y.com;RRHH\n")
        self.assertEqual(import_addresses(path), ["ana@x.com", "luis@y.com"])

    def test_csv_separado_por_comas(self):
        path = self._write("lista.csv", "nombre,correo\nAna,ana@x.com\n")
        self.assertEqual(import_addresses(path), ["ana@x.com"])

    def test_archivo_en_cp1252(self):
        path = self._write("lista.txt", "Ana Pérez ana@x.com\n", encoding="cp1252")
        self.assertEqual(import_addresses(path), ["ana@x.com"])

    def test_duplicados_se_eliminan(self):
        path = self._write("lista.txt", "ana@x.com\nANA@X.COM\nluis@y.com\n")
        self.assertEqual(import_addresses(path), ["ana@x.com", "luis@y.com"])

    def test_archivo_sin_direcciones(self):
        path = self._write("lista.txt", "no hay nada\n")
        with self.assertRaises(FileImportError):
            import_addresses(path)

    def test_archivo_inexistente(self):
        with self.assertRaises(FileImportError):
            import_addresses(self.dir / "no-existe.csv")

    def test_xls_viejo_avisa(self):
        path = self._write("lista.xls", "loquesea")
        with self.assertRaises(FileImportError) as ctx:
            import_addresses(path)
        self.assertIn(".xlsx", str(ctx.exception))

    @unittest.skipIf(openpyxl is None, "openpyxl no instalado")
    def test_xlsx_varias_hojas(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Nombre", "Correo"])
        sheet.append(["Ana", "ana@x.com"])
        sheet.append([None, None])
        sheet.append(["Luis", "luis@y.com"])
        segunda = workbook.create_sheet("Otros")
        segunda.append(["eva@z.com"])
        path = self.dir / "lista.xlsx"
        workbook.save(path)
        self.assertEqual(import_addresses(path), ["ana@x.com", "luis@y.com", "eva@z.com"])

    @unittest.skipIf(openpyxl is None, "openpyxl no instalado")
    def test_xlsx_roto_avisa(self):
        path = self._write("roto.xlsx", "esto no es un excel")
        with self.assertRaises(FileImportError):
            import_addresses(path)


class TestPlanillaConColumnas(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _csv(self, name, content):
        path = self.dir / name
        path.write_text(content, encoding="utf-8")
        return path

    def test_encabezados_y_columnas(self):
        path = self._csv(
            "datos.csv",
            "Area,Supervisor,Subordinado\n"
            "Calidad,sup1@x.com,Juan Pérez\n"
            "Planta,sup2@x.com,Marta Gómez\n",
        )
        table = read_table(path)
        self.assertEqual(table.columns, ["AREA", "SUBORDINADO"])
        self.assertEqual(table.emails, ["sup1@x.com", "sup2@x.com"])
        self.assertEqual(table.column_values("SUBORDINADO"), ["Juan Pérez", "Marta Gómez"])
        self.assertEqual(table.email_header, "Supervisor")

    def test_la_columna_de_correos_se_detecta_por_contenido(self):
        # El correo está en la última columna y el encabezado no lo dice.
        path = self._csv("datos.csv", "Caso,Quien\nuno,a@x.com\ndos,b@x.com\n")
        table = read_table(path)
        self.assertEqual(table.emails, ["a@x.com", "b@x.com"])
        self.assertEqual(table.columns, ["CASO"])

    def test_supervisor_repetido_conserva_una_fila_por_subordinado(self):
        path = self._csv(
            "datos.csv",
            "Supervisor,Subordinado\nsup@x.com,Juan\nsup@x.com,Luis\n",
        )
        table = read_table(path)
        self.assertEqual(table.emails, ["sup@x.com", "sup@x.com"])
        self.assertEqual(table.column_values("SUBORDINADO"), ["Juan", "Luis"])

    def test_sin_encabezado_no_hay_columnas(self):
        path = self._csv("datos.csv", "a@x.com\nb@x.com\n")
        table = read_table(path)
        self.assertEqual(table.columns, [])
        self.assertEqual(table.emails, ["a@x.com", "b@x.com"])

    def test_columna_vacia_se_ignora(self):
        path = self._csv("datos.csv", "Correo,Sobra,Caso\na@x.com,,uno\nb@x.com,,dos\n")
        self.assertEqual(read_table(path).columns, ["CASO"])

    def test_encabezado_repetido_se_toma_una_sola_vez(self):
        path = self._csv("datos.csv", "Correo,Caso,Caso\na@x.com,uno,otro\n")
        table = read_table(path)
        self.assertEqual(table.columns, ["CASO"])
        self.assertEqual(table.column_values("CASO"), ["uno"])

    def test_fila_sin_correo_se_saltea_con_sus_valores(self):
        path = self._csv("datos.csv", "Correo,Caso\na@x.com,uno\n,huerfano\nb@x.com,dos\n")
        table = read_table(path)
        self.assertEqual(table.emails, ["a@x.com", "b@x.com"])
        self.assertEqual(table.column_values("CASO"), ["uno", "dos"])

    def test_valor_faltante_queda_vacio(self):
        path = self._csv("datos.csv", "Correo,Caso\na@x.com,uno\nb@x.com,\n")
        self.assertEqual(read_table(path).column_values("CASO"), ["uno", ""])

    def test_archivo_sin_correos(self):
        path = self._csv("datos.csv", "Nombre,Caso\nAna,uno\n")
        with self.assertRaises(FileImportError):
            read_table(path)

    @unittest.skipIf(openpyxl is None, "openpyxl no instalado")
    def test_excel_con_columnas(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Supervisor", "Subordinado", "Motivo"])
        sheet.append(["sup1@x.com", "Juan", "vencimiento"])
        sheet.append([None, None, None])
        sheet.append(["sup2@x.com", "Marta", "ausencia"])
        path = self.dir / "datos.xlsx"
        workbook.save(path)
        table = read_table(path)
        self.assertEqual(table.columns, ["SUBORDINADO", "MOTIVO"])
        self.assertEqual(table.emails, ["sup1@x.com", "sup2@x.com"])
        self.assertEqual(table.column_values("MOTIVO"), ["vencimiento", "ausencia"])


if __name__ == "__main__":
    unittest.main()
