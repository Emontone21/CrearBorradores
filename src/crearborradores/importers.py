"""Importacion de destinatarios desde archivos Excel, CSV o TXT."""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from pathlib import Path

from .emails import EMAIL_RE, extract_addresses, is_valid_address
from .variables import is_valid_name, normalize_name

TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".log"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | EXCEL_EXTENSIONS

# Tipos para el diálogo "Abrir archivo" de la interfaz.
FILE_TYPES = [
    ("Todos los soportados", "*.xlsx *.xlsm *.csv *.tsv *.txt"),
    ("Excel", "*.xlsx *.xlsm *.xltx *.xltm"),
    ("CSV / texto", "*.csv *.tsv *.txt"),
    ("Todos los archivos", "*.*"),
]

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


class FileImportError(Exception):
    """Error legible para mostrarle al usuario."""


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in _ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    # latin-1 nunca falla, pero por las dudas:
    return raw.decode("latin-1", errors="replace")


def _read_csv(path: Path) -> str:
    """Devuelve el contenido del CSV/TSV como texto plano, celda por celda."""
    text = _read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    buffer = io.StringIO()
    for row in csv.reader(io.StringIO(text), dialect):
        buffer.write(" ".join(cell for cell in row if cell))
        buffer.write("\n")
    return buffer.getvalue()


def _read_excel(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - depende del entorno
        raise FileImportError(
            "Para leer archivos de Excel falta la librería 'openpyxl'.\n\n"
            "Instalala con:  pip install openpyxl\n"
            "O guarda la lista como .csv y volvé a intentar."
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise FileImportError(
            f"No se pudo abrir el archivo de Excel.\n\nDetalle: {exc}"
        ) from exc

    buffer = io.StringIO()
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is None:
                        continue
                    buffer.write(str(value))
                    buffer.write("\n")
    finally:
        workbook.close()
    return buffer.getvalue()


def import_addresses(path: str | os.PathLike[str]) -> list[str]:
    """Extrae todas las direcciones de correo de un archivo.

    Recorre todas las hojas/columnas y devuelve las direcciones encontradas
    sin duplicados, en el orden en que aparecen. Como el resultado se vuelca
    en el cuadro de destinatarios, el usuario puede revisarlo y editarlo
    antes de crear nada.
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise FileImportError(f"No se encontró el archivo:\n{file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".xls":
        raise FileImportError(
            "El formato .xls (Excel 97-2003) no está soportado.\n\n"
            "Abrilo en Excel y guardalo como .xlsx, o exportalo a .csv."
        )

    if suffix in EXCEL_EXTENSIONS:
        content = _read_excel(file_path)
    elif suffix in {".csv", ".tsv"}:
        content = _read_csv(file_path)
    else:
        content = _read_text_file(file_path)

    addresses = extract_addresses(content)
    if not addresses:
        raise FileImportError(
            f"No se encontró ninguna dirección de correo en:\n{file_path.name}"
        )
    return addresses


# ------------------------------------------------- planillas con columnas


@dataclass
class TableRow:
    """Una fila de la planilla: un correo y los valores de sus columnas."""

    email: str
    values: dict[str, str] = field(default_factory=dict)


@dataclass
class Table:
    """Planilla con una columna de correos y otras que pasan a ser variables."""

    columns: list[str] = field(default_factory=list)
    rows: list[TableRow] = field(default_factory=list)
    email_header: str = ""

    @property
    def emails(self) -> list[str]:
        return [row.email for row in self.rows]

    def column_values(self, name: str) -> list[str]:
        """Los valores de una columna, en el mismo orden que los correos."""
        return [row.values.get(name, "") for row in self.rows]


def _first_address(text: str) -> str:
    for candidate in EMAIL_RE.findall(text or ""):
        cleaned = candidate.strip().strip(".")
        if is_valid_address(cleaned):
            return cleaned
    return ""


def _matrix_from_excel(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - depende del entorno
        raise FileImportError(
            "Para leer archivos de Excel falta la librería 'openpyxl'.\n\n"
            "Instalala con:  pip install openpyxl\n"
            "O guarda la lista como .csv y volvé a intentar."
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise FileImportError(
            f"No se pudo abrir el archivo de Excel.\n\nDetalle: {exc}"
        ) from exc

    try:
        # Solo la hoja activa: una tabla con encabezados vive en una hoja.
        sheet = workbook.active
        matrix = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()
    return matrix


def _matrix_from_text(path: Path) -> list[list[str]]:
    text = _read_text_file(path)
    if path.suffix.lower() in {".csv", ".tsv"}:
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [
            [cell.strip() for cell in row]
            for row in csv.reader(io.StringIO(text), dialect)
        ]
    return [[line.strip()] for line in text.splitlines()]


def read_table(path: str | os.PathLike[str]) -> Table:
    """Lee una planilla como tabla: correos más una columna por variable.

    Toma como encabezado la primera fila, salvo que ya contenga una
    dirección de correo (en ese caso se entiende que son datos sueltos sin
    encabezado). La columna de correos se detecta por contenido, no por el
    nombre, así que da igual cómo se llame o en qué posición esté.

    Si la planilla tiene una sola columna, devuelve la tabla sin variables.
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise FileImportError(f"No se encontró el archivo:\n{file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".xls":
        raise FileImportError(
            "El formato .xls (Excel 97-2003) no está soportado.\n\n"
            "Abrilo en Excel y guardalo como .xlsx, o exportalo a .csv."
        )

    matrix = _matrix_from_excel(file_path) if suffix in EXCEL_EXTENSIONS else _matrix_from_text(file_path)
    matrix = [row for row in matrix if any(cell for cell in row)]
    if not matrix:
        raise FileImportError(f"El archivo está vacío:\n{file_path.name}")

    width = max(len(row) for row in matrix)
    matrix = [list(row) + [""] * (width - len(row)) for row in matrix]

    # La columna de correos es la que más direcciones válidas tenga.
    counts = [
        sum(1 for row in matrix if _first_address(row[index])) for index in range(width)
    ]
    email_index = max(range(width), key=lambda index: counts[index])
    if counts[email_index] == 0:
        raise FileImportError(
            f"No se encontró ninguna dirección de correo en:\n{file_path.name}"
        )

    has_header = not any(_first_address(cell) for cell in matrix[0])
    header = matrix[0] if has_header else [""] * width
    data = matrix[1:] if has_header else matrix

    columns: list[str] = []
    indexes: dict[str, int] = {}
    for index in range(width):
        if index == email_index:
            continue
        name = normalize_name(header[index])
        if not name or not is_valid_name(name) or name in indexes:
            continue
        if not any(row[index].strip() for row in data):
            continue  # columna vacía, no aporta nada
        columns.append(name)
        indexes[name] = index

    rows: list[TableRow] = []
    for row in data:
        email = _first_address(row[email_index])
        if not email:
            continue
        rows.append(
            TableRow(
                email=email,
                values={name: row[index].strip() for name, index in indexes.items()},
            )
        )

    if not rows:
        raise FileImportError(
            f"No se encontró ninguna dirección de correo en:\n{file_path.name}"
        )

    return Table(
        columns=columns,
        rows=rows,
        email_header=header[email_index] if has_header else "",
    )
